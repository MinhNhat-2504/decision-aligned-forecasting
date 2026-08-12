"""Export PAPER_MANUSCRIPT.md to a review-friendly Word document.

This intentionally supports only the Markdown constructs used by the manuscript:
headings, paragraphs, lists, pipe tables, display equations, and local figures.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PAPER_MANUSCRIPT.md"
TARGET = ROOT / "PAPER_MANUSCRIPT.docx"


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("*", "").replace("`", "")
    text = text.replace(r"\(", "").replace(r"\)", "")
    return text


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    add_page_number(sec.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04

    for name, size, color in (
        ("Title", 18, "17365D"),
        ("Heading 1", 14, "17365D"),
        ("Heading 2", 11.5, "2F5597"),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    if "Display Equation" not in doc.styles:
        eq = doc.styles.add_style("Display Equation", WD_STYLE_TYPE.PARAGRAPH)
        eq.font.name = "Cambria Math"
        eq.font.size = Pt(10)
        eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        eq.paragraph_format.space_before = Pt(3)
        eq.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Number"):
        doc.styles[style_name].font.name = "Times New Roman"
        doc.styles[style_name].font.size = Pt(10)


def add_pipe_table(doc: Document, block: list[str]) -> None:
    rows = [[clean_inline(x.strip()) for x in line.strip().strip("|").split("|")]
            for line in block]
    rows = [row for i, row in enumerate(rows) if i != 1]
    n_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Shading Accent 1"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "2F5597")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def export() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    title_parts: list[str] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("# ") and not line.startswith("## "):
            title_parts.append(clean_inline(line[2:]))
            i += 1
            continue
        if title_parts:
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(" ".join(title_parts)).bold = True
            title_parts.clear()

        if line == "---":
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2)
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if match:
                caption, relative = match.groups()
                path = ROOT / relative
                if path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(path), width=Inches(6.45))
                    cap = doc.add_paragraph(clean_inline(caption))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in cap.runs:
                        run.italic = True
                        run.font.size = Pt(8)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_pipe_table(doc, block)
            continue
        if line == r"\[":
            equation = []
            i += 1
            while i < len(lines) and lines[i].strip() != r"\]":
                equation.append(lines[i].strip())
                i += 1
            i += 1
            doc.add_paragraph(" ".join(equation), style="Display Equation")
            continue
        numbered = re.match(r"^\d+\.\s+(.*)", line)
        if numbered:
            item = [numbered.group(1)]
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if (not nxt or nxt.startswith("#") or nxt.startswith("|")
                        or nxt.startswith("![") or nxt == r"\["
                        or re.match(r"^\d+\.\s+", nxt) or nxt.startswith("- ")):
                    break
                item.append(nxt)
                i += 1
            doc.add_paragraph(clean_inline(" ".join(item)), style="List Number")
            continue
        if line.startswith("- "):
            item = [line[2:]]
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if (not nxt or nxt.startswith("#") or nxt.startswith("|")
                        or nxt.startswith("![") or nxt == r"\["
                        or re.match(r"^\d+\.\s+", nxt) or nxt.startswith("- ")):
                    break
                item.append(nxt)
                i += 1
            doc.add_paragraph(clean_inline(" ".join(item)), style="List Bullet")
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt == "---" or nxt.startswith("#")
                    or nxt.startswith("|") or nxt.startswith("![")
                    or nxt == r"\[" or re.match(r"^\d+\.\s+", nxt)
                    or nxt.startswith("- ")):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph(clean_inline(" ".join(paragraph_lines)))
        if p.text.startswith("Keywords:"):
            for run in p.runs:
                run.italic = True

    doc.core_properties.title = (
        "Forecast Value Is a Property of the Decision: "
        "Calibrated Quantile Demand Forecasting"
    )
    doc.core_properties.subject = "Working research manuscript"
    doc.core_properties.keywords = (
        "probabilistic forecasting, graph neural networks, conformal prediction, "
        "newsvendor, last-mile decision support"
    )
    doc.save(TARGET)
    print(f"Saved {TARGET}")


if __name__ == "__main__":
    export()
